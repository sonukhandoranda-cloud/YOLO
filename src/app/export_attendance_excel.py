from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.database.db import connect, init_db


HEADERS = [
    "Date",
    "Session",
    "Roll No",
    "Name",
    "Branch",
    "Batch",
    "Status",
    "Face Score",
    "Presence Score",
    "Timestamp",
]


def autosize_columns(sheet) -> None:
    for column_cells in sheet.columns:
        max_length = 0
        column = column_cells[0].column
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        sheet.column_dimensions[get_column_letter(column)].width = min(max_length + 2, 40)


def main() -> None:
    parser = argparse.ArgumentParser(description="Export attendance records to an Excel sheet.")
    parser.add_argument("--db", default="data/attendance/attendance.sqlite3")
    parser.add_argument("--date", default=date.today().isoformat(), help="Attendance date in YYYY-MM-DD format.")
    parser.add_argument("--session", help="Optional session name filter.")
    parser.add_argument("--output-dir", default="data/processed/attendance_exports")
    args = parser.parse_args()

    init_db(args.db)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_session = args.session.replace(" ", "_") if args.session else "all_sessions"
    output_path = output_dir / f"attendance_{args.date}_{safe_session}.xlsx"

    query = """
        SELECT
            a.session_date,
            a.name AS session_name,
            s.roll_no,
            s.name,
            s.branch,
            s.batch,
            r.decision_status,
            r.face_score,
            r.presence_score,
            r.timestamp
        FROM attendance_records r
        JOIN students s ON s.student_id = r.student_id
        JOIN attendance_sessions a ON a.session_id = r.session_id
        WHERE a.session_date = ?
    """
    params: list[str] = [args.date]
    if args.session:
        query += " AND a.name = ?"
        params.append(args.session)
    query += " ORDER BY a.name, s.roll_no, r.timestamp"

    conn = connect(args.db)
    try:
        rows = conn.execute(query, params).fetchall()
    finally:
        conn.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"
    sheet.append(HEADERS)

    for row in rows:
        sheet.append(
            [
                row["session_date"],
                row["session_name"],
                row["roll_no"],
                row["name"],
                row["branch"],
                row["batch"],
                row["decision_status"],
                round(float(row["face_score"]), 4),
                round(float(row["presence_score"]), 4),
                row["timestamp"],
            ]
        )

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    sheet.freeze_panes = "A2"
    autosize_columns(sheet)
    workbook.save(output_path)

    print(f"Exported {len(rows)} attendance records")
    print(f"Excel file: {output_path}")


if __name__ == "__main__":
    main()
